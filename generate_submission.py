import json
import uuid
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
DOCX_PATH = ROOT / "BaiTap5_API_QuanLyThuVien.docx"
COLLECTION_PATH = ROOT / "Library_API_Mock_Server.postman_collection.json"


BOOKS = [
    {
        "id": 1,
        "title": "Cho toi xin mot ve di tuoi tho",
        "author": "Nguyen Nhat Anh",
        "year": 2008,
        "available": True,
    },
    {
        "id": 2,
        "title": "Mat biec",
        "author": "Nguyen Nhat Anh",
        "year": 1990,
        "available": False,
    },
    {
        "id": 3,
        "title": "Dac nhan tam",
        "author": "Dale Carnegie",
        "year": 1936,
        "available": True,
    },
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_fixed(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_grid = table._tbl.tblGrid
    for grid_col in list(tbl_grid):
        tbl_grid.remove(grid_col)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_paragraph_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_run_font(run, name="Calibri", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_code_block(doc, text):
    for line in text.splitlines():
        p = doc.add_paragraph()
        p.style = "Code Block"
        p.paragraph_format.left_indent = Inches(0.12)
        p.paragraph_format.right_indent = Inches(0.08)
        p.paragraph_format.space_after = Pt(0)
        set_paragraph_shading(p, "F7F7F7")
        r = p.add_run(line if line else " ")
        set_run_font(r, "Courier New", 8.5)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def add_kv_table(doc, rows, widths=(2100, 7260)):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        set_cell_shading(cells[0], "F2F4F7")
        for cell in cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    set_run_font(run, size=9)
            cells[0].paragraphs[0].runs[0].bold = True
    set_table_fixed(table, list(widths))
    doc.add_paragraph()
    return table


def add_endpoint_table(doc):
    headers = ["Chuc nang", "Method", "URL", "Thanh cong", "Loi co the xay ra"]
    rows = [
        [
            "Lay danh sach tat ca sach",
            "GET",
            "/books",
            "200 OK",
            "500 Internal Server Error: loi xu ly tren server. 503 Service Unavailable: mock/API tam thoi khong san sang.",
        ],
        [
            "Lay chi tiet sach theo id",
            "GET",
            "/books/{id}",
            "200 OK",
            "400 Bad Request: id khong dung dinh dang so. 404 Not Found: khong ton tai sach voi id nay.",
        ],
        [
            "Them sach moi",
            "POST",
            "/books",
            "201 Created",
            "400 Bad Request: thieu title/author/year/available hoac sai kieu du lieu. 409 Conflict: id hoac ban ghi sach bi trung.",
        ],
        [
            "Cap nhat toan bo sach theo id",
            "PUT",
            "/books/{id}",
            "200 OK",
            "400 Bad Request: body khong day du truong bat buoc hoac id trong body lech URL. 404 Not Found: khong tim thay sach can cap nhat.",
        ],
        [
            "Xoa sach theo id",
            "DELETE",
            "/books/{id}",
            "204 No Content",
            "400 Bad Request: id khong hop le. 404 Not Found: khong co sach can xoa. 409 Conflict: sach dang duoc muon nen khong cho xoa.",
        ],
        [
            "Tim sach theo tac gia",
            "GET",
            "/books?author={author}",
            "200 OK",
            "400 Bad Request: author rong hoac query parameter sai dinh dang. 500 Internal Server Error: loi loc du lieu tren server.",
        ],
    ]
    widths = [1780, 920, 1740, 1180, 3740]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_fixed(table, widths)
    set_repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = header
        set_cell_shading(cell, "F2F4F7")
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    set_table_fixed(table, widths)
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                if col_idx in (1, 3):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    set_run_font(run, size=8.5, bold=(row_idx == 0))
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    doc.add_paragraph()


def add_test_matrix(doc):
    headers = ["Request mau", "Status", "Response mong doi"]
    rows = [
        ("GET {{baseUrl}}/books", "200", "Mang JSON gom 3 sach mau."),
        ("GET {{baseUrl}}/books/1", "200", "Mot object sach co id = 1."),
        ("POST {{baseUrl}}/books", "201", "Object sach moi, co id duoc tao va header Location."),
        ("PUT {{baseUrl}}/books/1", "200", "Object sach sau khi cap nhat toan bo."),
        ("DELETE {{baseUrl}}/books/1", "204", "Khong co body response."),
        ("GET {{baseUrl}}/books?author=Nguyen%20Nhat%20Anh", "200", "Mang JSON co cac sach cua Nguyen Nhat Anh."),
        ("GET {{baseUrl}}/books?author=Tac%20Gia%20Khong%20Ton%20Tai", "200", "Mang rong []. Tim khong co ket qua khong phai loi."),
    ]
    widths = [4320, 1040, 4000]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_table_fixed(table, widths)
    set_repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = header
        set_cell_shading(cell, "F2F4F7")
    for req, status, expected in rows:
        cells = table.add_row().cells
        cells[0].text = req
        cells[1].text = status
        cells[2].text = expected
    set_table_fixed(table, widths)
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                if col_idx == 1:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    set_run_font(run, size=8.8, bold=(row_idx == 0))
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    doc.add_paragraph()


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string("0B2545")
    title.paragraph_format.space_after = Pt(8)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    subtitle.font.size = Pt(12)
    subtitle.font.color.rgb = RGBColor.from_string("555555")
    subtitle.paragraph_format.space_after = Pt(14)

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    code_style = styles.add_style("Code Block", 1)
    code_style.font.name = "Courier New"
    code_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
    code_style.font.size = Pt(8.5)
    code_style.paragraph_format.space_after = Pt(0)
    code_style.paragraph_format.line_spacing = 1.0

    footer = section.footer.paragraphs[0]
    footer.text = "Bai tap 5 - API quan ly thu vien"
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in footer.runs:
        set_run_font(run, size=9, color="666666")


def build_docx():
    doc = Document()
    configure_document(doc)

    p = doc.add_paragraph(style="Title")
    p.add_run("Bài tập 5: Thiết kế API quản lý thư viện")
    p = doc.add_paragraph(style="Subtitle")
    p.add_run("Phần lý thuyết + cấu hình Postman Mock Server")

    doc.add_heading("1. Quy ước thiết kế", level=1)
    add_kv_table(
        doc,
        [
            ("Base URL", "{{baseUrl}} - thay bằng URL Mock Server do Postman cấp"),
            ("Resource chính", "/books"),
            ("Content-Type", "application/json"),
            ("Mô hình Book", "id, title, author, year, available"),
            ("Quy ước tìm kiếm", "Không có kết quả vẫn trả 200 OK với mảng rỗng [], không trả 404."),
        ],
    )

    doc.add_heading("2. Thiết kế endpoint", level=1)
    add_endpoint_table(doc)

    doc.add_heading("3. Ví dụ request/response JSON", level=1)
    doc.add_heading("3.1. Thêm một sách mới thành công", level=2)
    doc.add_paragraph("Request:")
    add_code_block(
        doc,
        """POST {{baseUrl}}/books
Content-Type: application/json

{
  "title": "Lap trinh Java co ban",
  "author": "Nguyen Van A",
  "year": 2024,
  "available": true
}""",
    )
    doc.add_paragraph("Response:")
    add_code_block(
        doc,
        """HTTP/1.1 201 Created
Location: /books/4
Content-Type: application/json

{
  "id": 4,
  "title": "Lap trinh Java co ban",
  "author": "Nguyen Van A",
  "year": 2024,
  "available": true
}""",
    )

    doc.add_heading("3.2. Tìm sách theo tác giả không có kết quả", level=2)
    doc.add_paragraph("Request:")
    add_code_block(
        doc,
        """GET {{baseUrl}}/books?author=Tac%20Gia%20Khong%20Ton%20Tai
Accept: application/json""",
    )
    doc.add_paragraph("Response:")
    add_code_block(
        doc,
        """HTTP/1.1 200 OK
Content-Type: application/json

[]""",
    )

    doc.add_heading("4. Thực hành trên Postman Mock Server", level=1)
    doc.add_paragraph("Các bước thực hiện trong Postman:", style="Normal")
    for step in [
        "Import file Library_API_Mock_Server.postman_collection.json.",
        "Chọn collection vừa import, tạo Mock Server từ collection.",
        "Copy Mock Server URL Postman cấp và cập nhật biến collection baseUrl.",
        "Gửi thử từng request trong collection, kiểm tra status code và JSON trả về.",
        "Chụp màn hình kết quả test để chèn vào file báo cáo nếu giảng viên yêu cầu minh chứng.",
    ]:
        doc.add_paragraph(step, style="List Number")
    doc.add_heading("Ma trận kiểm thử dự kiến", level=2)
    add_test_matrix(doc)
    doc.add_paragraph(
        "Gợi ý ảnh chụp cần bổ sung sau khi chạy Postman: màn hình Mock Server đã tạo, "
        "GET /books, GET /books/1, POST /books, PUT /books/1, DELETE /books/1 và "
        "GET /books?author=Tac%20Gia%20Khong%20Ton%20Tai trả về []."
    )

    doc.add_heading("5. Lý thuyết: PUT và PATCH", level=1)
    doc.add_paragraph(
        "PUT dùng để thay thế toàn bộ tài nguyên tại URL đã chỉ định. Khi client gửi PUT /books/{id}, "
        "body nên chứa đầy đủ thông tin của sách. Nếu thiếu trường bắt buộc, server có thể trả 400 Bad Request. "
        "PUT thường có tính idempotent: gửi cùng một request nhiều lần vẫn cho cùng trạng thái tài nguyên."
    )
    doc.add_paragraph(
        "PATCH dùng để cập nhật một phần tài nguyên. Client chỉ gửi các trường cần thay đổi, ví dụ chỉ đổi "
        "available từ true sang false. PATCH phù hợp khi muốn sửa một vài thuộc tính mà không gửi lại toàn bộ object."
    )
    doc.add_paragraph(
        "Trong bài này yêu cầu là 'cập nhật toàn bộ thông tin sách theo id', vì vậy nên dùng PUT. "
        "Nếu bài toán bổ sung chức năng cập nhật một phần như chỉ đổi trạng thái mượn/trả sách, khi đó nên thiết kế thêm PATCH."
    )

    doc.add_heading("6. File nộp kèm", level=1)
    doc.add_paragraph("File Word/PDF lý thuyết: BaiTap5_API_QuanLyThuVien.docx hoặc PDF xuất từ file này.")
    doc.add_paragraph("File export Postman collection: Library_API_Mock_Server.postman_collection.json.")

    doc.save(DOCX_PATH)


def response_example(name, original_request, code, status, body, headers=None):
    return {
        "name": name,
        "originalRequest": original_request,
        "status": status,
        "code": code,
        "_postman_previewlanguage": "json",
        "header": headers or [{"key": "Content-Type", "value": "application/json"}],
        "cookie": [],
        "body": body,
    }


def url(path, query=None):
    raw = "{{baseUrl}}" + path
    query_items = []
    if query:
        raw += "?" + "&".join(f"{k}={v.replace(' ', '%20')}" for k, v in query.items())
        query_items = [{"key": k, "value": v} for k, v in query.items()]
    data = {"raw": raw, "host": ["{{baseUrl}}"], "path": [p for p in path.split("/") if p]}
    if query_items:
        data["query"] = query_items
    return data


def request(method, path, query=None, body=None):
    req = {
        "method": method,
        "header": [{"key": "Accept", "value": "application/json"}],
        "url": url(path, query),
    }
    if body is not None:
        req["header"].append({"key": "Content-Type", "value": "application/json"})
        req["body"] = {"mode": "raw", "raw": json.dumps(body, ensure_ascii=False, indent=2), "options": {"raw": {"language": "json"}}}
    return req


def item(name, method, path, query=None, body=None, responses=None):
    req = request(method, path, query, body)
    return {"name": name, "request": req, "response": responses or []}


def build_collection():
    post_body = {
        "title": "Lap trinh Java co ban",
        "author": "Nguyen Van A",
        "year": 2024,
        "available": True,
    }
    created = {"id": 4, **post_body}

    put_body = {
        "id": 1,
        "title": "Cho toi xin mot ve di tuoi tho",
        "author": "Nguyen Nhat Anh",
        "year": 2008,
        "available": False,
    }

    collection_items = []

    req = request("GET", "/books")
    collection_items.append(
        {
            "name": "GET - Lay danh sach tat ca sach",
            "request": req,
            "response": [
                response_example(
                    "200 OK - Books list",
                    req,
                    200,
                    "OK",
                    json.dumps(BOOKS, ensure_ascii=False, indent=2),
                )
            ],
        }
    )

    req = request("GET", "/books/1")
    collection_items.append(
        {
            "name": "GET - Lay chi tiet sach theo id",
            "request": req,
            "response": [
                response_example(
                    "200 OK - Book detail",
                    req,
                    200,
                    "OK",
                    json.dumps(BOOKS[0], ensure_ascii=False, indent=2),
                )
            ],
        }
    )

    req = request("POST", "/books", body=post_body)
    collection_items.append(
        {
            "name": "POST - Them sach moi",
            "request": req,
            "response": [
                response_example(
                    "201 Created - Book created",
                    req,
                    201,
                    "Created",
                    json.dumps(created, ensure_ascii=False, indent=2),
                    headers=[
                        {"key": "Content-Type", "value": "application/json"},
                        {"key": "Location", "value": "/books/4"},
                    ],
                )
            ],
        }
    )

    req = request("PUT", "/books/1", body=put_body)
    collection_items.append(
        {
            "name": "PUT - Cap nhat toan bo sach theo id",
            "request": req,
            "response": [
                response_example(
                    "200 OK - Book replaced",
                    req,
                    200,
                    "OK",
                    json.dumps(put_body, ensure_ascii=False, indent=2),
                )
            ],
        }
    )

    req = request("DELETE", "/books/1")
    collection_items.append(
        {
            "name": "DELETE - Xoa sach theo id",
            "request": req,
            "response": [
                response_example(
                    "204 No Content - Book deleted",
                    req,
                    204,
                    "No Content",
                    "",
                    headers=[],
                )
            ],
        }
    )

    req = request("GET", "/books", query={"author": "Nguyen Nhat Anh"})
    collection_items.append(
        {
            "name": "GET - Tim sach theo tac gia co ket qua",
            "request": req,
            "response": [
                response_example(
                    "200 OK - Books by author",
                    req,
                    200,
                    "OK",
                    json.dumps([BOOKS[0], BOOKS[1]], ensure_ascii=False, indent=2),
                )
            ],
        }
    )

    req = request("GET", "/books", query={"author": "Tac Gia Khong Ton Tai"})
    collection_items.append(
        {
            "name": "GET - Tim sach theo tac gia khong co ket qua",
            "request": req,
            "response": [
                response_example(
                    "200 OK - Empty result",
                    req,
                    200,
                    "OK",
                    "[]",
                )
            ],
        }
    )

    collection = {
        "info": {
            "_postman_id": str(uuid.uuid4()),
            "name": "BaiTap5 - Library API Mock Server",
            "description": "Collection mau cho bai tap thiet ke API quan ly thu vien va tao Postman Mock Server.",
            "schema": "https://schema.getpostman.com/json/collection/v2.1.0/collection.json",
        },
        "item": collection_items,
        "variable": [
            {
                "key": "baseUrl",
                "value": "https://<your-mock-id>.mock.pstmn.io",
                "type": "string",
            }
        ],
    }
    COLLECTION_PATH.write_text(json.dumps(collection, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    build_docx()
    build_collection()
    print(f"Created {DOCX_PATH}")
    print(f"Created {COLLECTION_PATH}")
